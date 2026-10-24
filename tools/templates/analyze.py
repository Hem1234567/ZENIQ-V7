import sys, re, json, collections
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def walk(shapes, ox=0, oy=0, sx=1.0, sy=1.0):
    """yield (shape, abs_left, abs_top, abs_w, abs_h) in EMU, applying group transforms"""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            x = sh._element.grpSpPr.xfrm
            cw = x.chExt.cx or 1; ch = x.chExt.cy or 1
            nsx = sx * x.ext.cx / cw; nsy = sy * x.ext.cy / ch
            # child abs = group.off(abs) + (child - chOff)*scale
            gl = ox + (x.off.x - 0) * sx if False else None
            gx = ox + x.off.x * sx; gy = oy + x.off.y * sy
            # children are positioned relative to chOff
            yield from walk_children(sh, gx, gy, x.chOff.x, x.chOff.y, nsx, nsy)
        else:
            yield sh, ox + sh.left * sx, oy + sh.top * sy, sh.width * sx, sh.height * sy

def walk_children(group, gx, gy, cox, coy, sx, sy):
    for sh in group.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            x = sh._element.grpSpPr.xfrm
            cw = x.chExt.cx or 1; ch = x.chExt.cy or 1
            ngx = gx + (x.off.x - cox) * sx; ngy = gy + (x.off.y - coy) * sy
            yield from walk_children(sh, ngx, ngy, x.chOff.x, x.chOff.y, sx * x.ext.cx / cw, sy * x.ext.cy / ch)
        else:
            yield sh, gx + (sh.left - cox) * sx, gy + (sh.top - coy) * sy, sh.width * sx, sh.height * sy

def fsize(sh):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size: return r.font.size.pt
    return 18.0

def month_re(t):
    return re.search(r'\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+\d{0,2},?\s*\d{4}\b|\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b|\b\d{1,2}\s+(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4}\b', t, re.I)

def analyze(path, keep):
    prs = Presentation(path)
    slides = list(prs.slides)
    # collect text shapes
    info = {}
    freq = collections.Counter()
    for i, s in enumerate(slides, 1):
        L = []
        for sh, l, t, w, h in walk(s.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = sh.text_frame.text.strip()
                L.append(dict(id=sh.shape_id, l=l/12700, t=t/12700, w=w/12700, h=h/12700, f=fsize(sh), txt=txt, np=len(sh.text_frame.paragraphs)))
        info[i] = L
        for tx in set(x['txt'] for x in L): freq[tx] += 1
    out = []
    for i in keep:
        L = info[i]
        decor, num, content = [], [], []
        for x in L:
            tx = x['txt']
            if freq[tx] >= 3 and len(tx) < 40 and not tx.lower().startswith('lorem'): decor.append(x)
            elif re.fullmatch(r'\d{1,3}\.?', tx): num.append(x)
            else: content.append(x)
        out.append(dict(i=i, decor=decor, num=num, content=content))
    return out

if __name__ == '__main__':
    path = sys.argv[1]; keep = [int(x) for x in sys.argv[2].split(',')]
    for r in analyze(path, keep):
        print('--- slide', r['i'], 'decor:', [d['txt'][:14] for d in r['decor']], 'num:', [n['txt'] for n in r['num']])
        for x in sorted(r['content'], key=lambda z: (round(z['t']/20), z['l'])):
            print('   id=%-3d l=%4.0f t=%4.0f %4.0fx%-4.0f f=%-5.0f np=%d |%s' % (x['id'], x['l'], x['t'], x['w'], x['h'], x['f'], x['np'], x['txt'][:40].replace('\n', '/')))
